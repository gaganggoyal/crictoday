"""Writes the .xlsx files tests/domain/xlsx.test.ts reads.

    python3 -m venv /tmp/xlsx && /tmp/xlsx/bin/pip install openpyxl==3.1.5
    /tmp/xlsx/bin/python tests/fixtures/make-xlsx.py tests/fixtures

openpyxl writes text inline. schedule-excel.xlsx moves it into a shared string table and
prefixes the sheet's elements, as Excel, Google Sheets and the .NET Open XML SDK save files.
"""
import datetime, sys, zipfile
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.utils.datetime import CALENDAR_MAC_1904

out = sys.argv[1]

def build(path, epoch=None):
    wb = Workbook()
    if epoch: wb.epoch = epoch
    ws = wb.active
    ws.title = "Fixtures"
    ws["A1"] = "Diwali Cup 2026 – fixtures & venues"
    ws.merge_cells("A1:F1")
    ws.append([])
    ws.append(["Date", "Start time", "Match", "Venue", "Town", "Entry", "Notes", "Confirmed"])
    ws.append([datetime.date(2026, 11, 1), datetime.time(9, 30), "Dadar Strikers vs Matunga Lions",
               "Shivaji Park", "Mumbai", "Free", "Gate 2 <north> & east", True])
    ws.append([datetime.datetime(2026, 11, 8, 14, 0), None, "Matunga Lions vs Pune Panthers",
               "Nehru Stadium", "Pune", "₹50 at the gate", None, False])
    # Dates typed as text, a rich-text cell, and a formula saved without its result.
    ws.append(["15/11/2026", "4:00 PM", None, "Dadoji Konddev", "Thane", "Free", "=1+1", None])
    ws["C6"].value = CellRichText([TextBlock(InlineFont(b=True), "Thane Tigers"), " vs Dadar Strikers"])
    ws.cell(row=8, column=1, value="Last row after a gap")
    other = wb.create_sheet("Notes")
    other["A1"] = "This sheet is not read"
    wb.save(path)

build(f"{out}/schedule.xlsx")
build(f"{out}/schedule-1904.xlsx", CALENDAR_MAC_1904)

# The same workbook with every part stored rather than deflated, as some tools write it.
with zipfile.ZipFile(f"{out}/schedule.xlsx") as source, zipfile.ZipFile(f"{out}/schedule-stored.xlsx", "w", zipfile.ZIP_STORED) as target:
    for item in source.infolist():
        target.writestr(item.filename, source.read(item.filename))


# ---- Shared strings, prefixes and a zip bomb, from schedule.xlsx ----
import re, struct

fixtures = out
with zipfile.ZipFile(f"{fixtures}/schedule.xlsx") as z:
    files = {i.filename: z.read(i.filename) for i in z.infolist()}

# Excel and Google Sheets keep text in a shared string table; turn the inline strings into one.
sheet = files["xl/worksheets/sheet1.xml"].decode()
strings = []
def shared(match):
    strings.append(match.group(2))
    return f'<c{match.group(1).replace(" t=\"inlineStr\"", " t=\"s\"")}><v>{len(strings) - 1}</v></c>'
sheet = re.sub(r'<c([^>]*t="inlineStr"[^>]*)><is>(.*?)</is></c>', shared, sheet)
# Some writers prefix every element, as the .NET Open XML SDK does.
sheet = sheet.replace('<worksheet xmlns=', '<x:worksheet xmlns:x=')
for tag in ["worksheet", "sheetPr", "outlinePr", "pageSetUpPr", "dimension", "sheetViews", "sheetView",
            "selection", "sheetFormatPr", "sheetData", "row", "c", "v", "f", "mergeCells", "mergeCell",
            "pageMargins"]:
    sheet = re.sub(r"<(/?)" + tag + r"\b", r"<\1x:" + tag, sheet)
items = []
for text in strings:
    if text == "<t>Shivaji Park</t>":
        # A phonetic guide, as Japanese Excel adds; it is not part of the text.
        text += '<rPh sb="0" eb="7"><t>シヴァージー</t></rPh>'
    if "Gate 2" in text:
        text = text.replace("Gate 2 ", "Gate 2_x000D_")
    items.append(f"<si>{text}</si>")
files["xl/worksheets/sheet1.xml"] = sheet.encode()
files["xl/sharedStrings.xml"] = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(items)}" uniqueCount="{len(items)}">'
    + "".join(items) + "</sst>").encode()
files["[Content_Types].xml"] = files["[Content_Types].xml"].replace(
    b"</Types>",
    b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml" /></Types>')
files["xl/_rels/workbook.xml.rels"] = files["xl/_rels/workbook.xml.rels"].replace(
    b"</Relationships>",
    b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml" Id="rId9" /></Relationships>')
with zipfile.ZipFile(f"{fixtures}/schedule-excel.xlsx", "w", zipfile.ZIP_DEFLATED) as target:
    for name, data in files.items():
        target.writestr(name, data)

# A workbook whose sheet claims to be 100 bytes but inflates past the reader's limit.
bomb = dict(files)
bomb["xl/worksheets/sheet1.xml"] = b"<x:worksheet>" + b"<x:row/>" * 3_000_000 + b"</x:worksheet>"
path = f"{fixtures}/schedule-bomb.xlsx"
with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as target:
    for name, data in bomb.items():
        target.writestr(name, data)
data = bytearray(open(path, "rb").read())
name = b"xl/worksheets/sheet1.xml"
at = 0
while True:
    at = data.index(b"PK\x01\x02", at)
    length = struct.unpack_from("<H", data, at + 28)[0]
    if data[at + 46 : at + 46 + length] == name:
        struct.pack_into("<I", data, at + 24, 100)
        break
    at += 4
open(path, "wb").write(bytes(data))
print(len(items), "shared strings")
